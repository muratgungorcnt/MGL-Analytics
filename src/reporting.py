"""
Excel report generator
"""

import logging
from pathlib import Path
from datetime import datetime
from typing import Dict, List
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from src.config import ConfigManager
from src.database import DatabaseManager
from src.decorators import safe_run

logger = logging.getLogger(__name__)


class ExcelReporter:
    """
    Generates Excel reports from database signals.
    """

    def __init__(self, config: ConfigManager, db: DatabaseManager):
        """
        Initialize Excel reporter.

        Args:
            config: Configuration manager
            db: Database manager
        """
        self.config = config
        self.db = db
        self.output_dir = Path("reports")
        self.output_dir.mkdir(exist_ok=True)

    @safe_run(default_return=False, log_level="ERROR")
    def generate_report(self, filename: str = None) -> bool:
        """
        Generate Excel report from database.

        Args:
            filename: Output filename (optional)

        Returns:
            True on success, False on error
        """
        if not filename:
            filename = self.config.get(
                "reports.excel_filename",
                "MGL_Analytics_Rapor.xlsx"
            )

        filepath = self.output_dir / filename

        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Market Analysis"

        # Set up headers
        headers = [
            "Symbol",
            "Type",
            "Timestamp",
            "Price",
            "RSI",
            "MACD",
            "Signal",
            "Stochastic",
            "Volume",
            "Volume Change %",
            "Direction",
            "Decision"
        ]
        ws.append(headers)

        # Style headers
        header_fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid"
        )
        header_font = Font(color="FFFFFF", bold=True, size=12)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Get data from database
        all_signals = self.db.get_all_signals(limit=500)

        if not all_signals:
            logger.warning("No signals to export")
            return False

        # Add data rows
        row_num = 2
        for symbol, signals in sorted(all_signals.items()):
            for signal in signals[-10:]:  # Last 10 per symbol
                try:
                    ws.append([
                        symbol,
                        signal.get("signal_type", "UNKNOWN"),
                        signal.get("timestamp", ""),
                        signal.get("price", 0),
                        signal.get("rsi", 0),
                        signal.get("macd", 0),
                        signal.get("macd_signal", 0),
                        signal.get("stochastic", 0),
                        signal.get("volume", 0),
                        signal.get("volume_change", 0),
                        signal.get("direction", ""),
                        signal.get("decision", "")
                    ])

                    # Color code RSI cells
                    rsi_val = signal.get("rsi", 50)
                    rsi_cell = ws[f"E{row_num}"]

                    if rsi_val >= 70:
                        rsi_cell.fill = PatternFill(
                            start_color="FF6B6B",
                            end_color="FF6B6B",
                            fill_type="solid"
                        )
                    elif rsi_val <= 30:
                        rsi_cell.fill = PatternFill(
                            start_color="4ECDC4",
                            end_color="4ECDC4",
                            fill_type="solid"
                        )

                    row_num += 1

                except Exception as e:
                    logger.warning(f"Row write error: {e}")
                    continue

        # Adjust column widths
        widths = [12, 12, 20, 12, 10, 12, 12, 12, 15, 15, 12, 20]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        # Save file
        wb.save(str(filepath))
        logger.info(f"✅ Excel report generated: {filepath}")
        return True

    def generate_daily_summary(self) -> bool:
        """
        Generate daily summary report.

        Returns:
            True on success
        """
        today = datetime.now().strftime("%Y-%m-%d")
        filename = f"MGL_Analytics_Summary_{today}.xlsx"
        return self.generate_report(filename)

    def __repr__(self) -> str:
        return f"ExcelReporter(output_dir={self.output_dir})"
